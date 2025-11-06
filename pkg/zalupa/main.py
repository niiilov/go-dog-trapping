from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.styles import Font, PatternFill, Border, Alignment, Protection
from pathlib import Path
from fastapi import FastAPI
from fastapi.responses import JSONResponse
from pydantic import BaseModel
from typing import List, Optional
from datetime import datetime
import uvicorn
import os
from typing import List, Dict, Any
from copy import copy

app = FastAPI()


@app.get("/")
async def root():
    return {"message": "API работает", "endpoints": ["/generate", "/generate-multiple"]}


@app.get("/health")
async def health():
    return {"status": "ok", "timestamp": datetime.now().isoformat()}


class RequestItem(BaseModel):
    source_name: str
    applicant_name: str
    address: str
    dogs_count: int
    behavior: str
    urgency: str
    contact_person: str


class RequestFull(BaseModel):
    number: str
    source_name: str
    applicant_name: str
    address: str
    dogs_count: int
    behavior: str
    urgency: str
    contact_person: str


class RequestMultiple(BaseModel):
    number: str
    requests: List[RequestItem]

class DogRow(BaseModel):
    source_name: str
    applicant_name: str
    address: str
    dogs_count: int
    behavior: str
    urgency: str
    contact_person: str



class RequestMultipleInsert(BaseModel):
    number: str
    rows: List[DogRow]
    start_row: Optional[int] = 20


def copy_images_to_sheet(source_sheet, target_sheet):
    """
    Копируем все изображения с исходного листа на целевой
    """
    try:
        from openpyxl.drawing.image import Image
        import io

        print(f"Копируем изображения с {source_sheet.title} на {target_sheet.title}")

        # Проверяем есть ли изображения на исходном листе
        if hasattr(source_sheet, "_images") and source_sheet._images:
            print(f"Найдено {len(source_sheet._images)} изображений")

            for img in source_sheet._images:
                try:
                    # Создаем новое изображение из данных оригинала
                    img_data = img.ref.getvalue()
                    new_img = Image(io.BytesIO(img_data))

                    # Копируем позицию и размер
                    new_img.anchor = img.anchor
                    new_img.width = img.width
                    new_img.height = img.height

                    # Добавляем на целевой лист
                    target_sheet.add_image(new_img)
                    print(f"✓ Изображение скопировано")

                except Exception as e:
                    print(f"✗ Ошибка копирования изображения: {e}")

        else:
            print("Изображений на исходном листе не найдено")

    except Exception as e:
        print(f"✗ Общая ошибка копирования изображений: {e}")


def generate_single_document(
    template_path: str, output_path: str, data: Dict[str, Any]
):
    """
    Генерация Excel-документа на основе шаблона для одной заявки
    """
    wb = load_workbook(template_path)
    ws = wb.active

    # Подставляем данные
    print(f"Всего данных для записи: {len(data)} элементов")
    for cell_address, value in data.items():
        print(f"Попытка записи в {cell_address} = {value} (тип: {type(value)})")
        try:
            # Находим все объединенные диапазоны
            merged_ranges = list(ws.merged_cells.ranges)
            target_cell = None

            # Проверяем, находится ли ячейка в объединенном диапазоне
            for merged_range in merged_ranges:
                if cell_address in merged_range:
                    min_col, min_row, max_col, max_row = merged_range.bounds
                    target_cell = ws.cell(row=min_row, column=min_col).coordinate
                    print(
                        f"Ячейка {cell_address} в объединенном диапазоне, используем {target_cell}"
                    )
                    break

            if target_cell:
                ws[target_cell] = value
                print(f"✓ Записано в объединенную ячейку {target_cell}: {value}")
            else:
                ws[cell_address] = value
                print(f"✓ Записано в ячейку {cell_address}: {value}")

        except Exception as e:
            print(f"✗ Ошибка при записи в ячейку {cell_address}: {e}")
            try:
                ws[cell_address].value = value
                print(f"✓ Записано напрямую в {cell_address}: {value}")
            except Exception as e2:
                print(f"✗ Не удалось записать в ячейку {cell_address}: {e2}")

    shared_dir = Path("/app/shared")
    wb.save(f"{shared_dir}/{output_path}")



@app.post("/generate")
async def generate_doc(req: RequestFull):
    try:
        print(f"Получен запрос: {req}")
        filename = f"zayavka_{req.number}_{datetime.now().year}.xlsx"
        template = "template.xlsx"

        if not os.path.exists(template):
            return JSONResponse(
                {"error": f"Template file {template} not found"}, status_code=500
            )

        data = {
            "B15": f"Заявка № {req.number}",
            "B12": datetime.now().strftime("%d.%m.%Y"),
            "B20": req.source_name,
            "D20": req.applicant_name,
            "E20": req.address,
            "F20": req.dogs_count,
            "G20": req.behavior,
            "H20": req.urgency,
            "I20": req.contact_person,
        }

        generate_single_document(template, filename, data)
        print("Документ успешно создан:", filename)
        shared_dir = Path("/app/shared")
        if os.path.exists(f"{shared_dir}/{filename}"):
            return JSONResponse({"status": "ok", "filename": filename}, status_code=200)
        return JSONResponse({"error": "File not created"}, status_code=500)
    except Exception as e:
        print(f"Ошибка при создании документа: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)





def generate_document_multiple(template_path, output_path, rows, start_row=20, header=None):
    wb = load_workbook(template_path)
    ws = wb.active

    if header:
        for cell, value in header.items():
            ws[cell] = value

    merged_to_copy = []
    if rows:
        # decide which row to use as style/template source.
        # If there are merged ranges that start at start_row, use start_row as template source,
        # otherwise prefer the row above (start_row-1) when available.
        all_merged = list(ws.merged_cells.ranges)
        has_merged_on_start = False
        for mr in all_merged:
            min_col, min_row, max_col_mr, max_row = range_boundaries(str(mr))
            if min_row == start_row:
                has_merged_on_start = True
                break

        template_row_idx = start_row if has_merged_on_start or start_row <= 1 else start_row - 1
        orig_max_col = ws.max_column or 1

        # Сохраняем стили и значения (формулы удаляем из шаблона)
        col_styles = {}
        template_values = {}
        for col in range(1, orig_max_col + 1):
            # for columns B..I (2..9) take styles specifically from row 20 (as requested)
            style_row = template_row_idx
            if 2 <= col <= 9:
                style_row = 20
            cell = ws.cell(row=style_row, column=col)
            col_styles[col] = {
                "font": copy(cell.font),
                "border": copy(cell.border),
                "fill": copy(cell.fill),
                "number_format": cell.number_format,
                "protection": copy(cell.protection),
                "alignment": copy(cell.alignment),
                "_style": copy(getattr(cell, "_style", None)),
            }
            val = cell.value
            if isinstance(val, str) and val.startswith("="):
                # удаляем формулу из шаблонной строки
                cell.value = None
                template_values[col] = None
            else:
                # Do not copy cell values from template — we only want formatting (user data will be written later)
                template_values[col] = None

        # Сохраним атрибуты RowDimension
        row_dim = ws.row_dimensions.get(template_row_idx)
        template_row_attrs = {
            "height": getattr(row_dim, "height", None),
            "hidden": getattr(row_dim, "hidden", False),
            "outline_level": getattr(row_dim, "outline_level", 0),
        }

        # Сохраним ширины колонок (если нужно)
        col_widths = {}
        for c in range(1, orig_max_col + 1):
            col_widths[c] = ws.column_dimensions[get_column_letter(c)].width

        # Собираем merged ranges, которые лежат на строке 20 и пересекают диапазон B..I
        for mr in list(ws.merged_cells.ranges):
            min_col, min_row, max_col_mr, max_row = range_boundaries(str(mr))
            if min_row == 20 and max_row == 20 and not (max_col_mr < 2 or min_col > 9):
                # ограничим копируемый диапазон колонок рамками объединения
                merged_to_copy.append((min_col, max_col_mr))

        # BEFORE inserting rows: save styles and row dimensions for all rows at or below start_row
        saved_rows = []
        max_row_before = ws.max_row
        # Also save row 22 specifically (full copy: values + styles + merged info)
        row22_idx = 22
        saved_row22 = {}
        saved_row22_cells = {}
        saved_row22_merged = []
        if 1 <= row22_idx <= max_row_before:
            for c in range(1, orig_max_col + 1):
                cell = ws.cell(row=row22_idx, column=c)
                saved_row22_cells[c] = {
                    "value": copy(cell.value),
                    "font": copy(cell.font),
                    "border": copy(cell.border),
                    "fill": copy(cell.fill),
                    "alignment": copy(cell.alignment),
                    "number_format": cell.number_format,
                    "protection": copy(cell.protection),
                    "_style": copy(getattr(cell, "_style", None)),
                }
            # merged ranges that cover row22
            for mr in list(ws.merged_cells.ranges):
                min_col, min_row, max_col_mr, max_row = range_boundaries(str(mr))
                if min_row <= row22_idx <= max_row:
                    saved_row22_merged.append((min_col, max_col_mr, min_row, max_row))
            saved_row22["cells"] = saved_row22_cells
            saved_row22["merged"] = saved_row22_merged
            # save row22 row dimension
            saved_row22["row_dim"] = ws.row_dimensions.get(row22_idx)
        for r in range(start_row, max_row_before + 1):
            row_dim = ws.row_dimensions.get(r)
            row_cells = {}
            for c in range(1, orig_max_col + 1):
                cell = ws.cell(row=r, column=c)
                row_cells[c] = {
                    "font": copy(cell.font),
                    "border": copy(cell.border),
                    "fill": copy(cell.fill),
                    "alignment": copy(cell.alignment),
                    "number_format": cell.number_format,
                    "protection": copy(cell.protection),
                    "_style": copy(getattr(cell, "_style", None)),
                }
            saved_rows.append((r, row_dim, row_cells))

        # Insert blank rows to push existing template content down so rows below remain unchanged
        ws.insert_rows(start_row, amount=len(rows))

    # After insertion, reapply saved styles to the shifted original rows
        shift = len(rows)
        for (orig_r, row_dim, row_cells) in saved_rows:
            new_r = orig_r + shift
            # restore row dimension attributes
            if row_dim is not None:
                rd = ws.row_dimensions[new_r]
                try:
                    if getattr(row_dim, "height", None) is not None:
                        rd.height = row_dim.height
                except Exception:
                    pass
                try:
                    rd.hidden = getattr(row_dim, "hidden", False)
                except Exception:
                    pass
                try:
                    rd.outline_level = getattr(row_dim, "outline_level", 0)
                except Exception:
                    pass
            # restore cell styles
            for c, style in row_cells.items():
                try:
                    target = ws.cell(row=new_r, column=c)
                    _apply_style(target, style)
                except Exception:
                    pass

        # If we saved row22, paste it to its new position (row22 + shift) and recreate merged ranges there
        if saved_row22:
            new_row22 = row22_idx + shift
            # paste cell values and styles
            for c, info in saved_row22["cells"].items():
                try:
                    target = ws.cell(row=new_row22, column=c)
                    # set value
                    target.value = copy(info.get("value"))
                    # apply styles
                    try:
                        _apply_style(target, info)
                    except Exception:
                        pass
                except Exception:
                    pass
            # recreate merged ranges that covered row22 (adjusted to new_row22)
            for min_col, max_col_mr, min_row, max_row in saved_row22["merged"]:
                # create a merged range only if it spans this single row, otherwise shift accordingly
                height = max_row - min_row
                if height == 0:
                    left = f"{get_column_letter(min_col)}{new_row22}"
                    right = f"{get_column_letter(max_col_mr)}{new_row22}"
                    try:
                        ws.merge_cells(f"{left}:{right}")
                    except Exception:
                        pass
                else:
                    # if merged spanned multiple rows, recreate the same height starting at new_row22 - (row22_idx - min_row)
                    offset = row22_idx - min_row
                    new_min_row = new_row22 - offset
                    new_max_row = new_min_row + (max_row - min_row)
                    left = f"{get_column_letter(min_col)}{new_min_row}"
                    right = f"{get_column_letter(max_col_mr)}{new_max_row}"
                    try:
                        ws.merge_cells(f"{left}:{right}")
                    except Exception:
                        pass

            # Finally, clear original row22 position (the old place) - clear values
            # and copy styles from the row above (row21) so it matches surrounding rows.
            try:
                for c in range(1, orig_max_col + 1):
                    cell = ws.cell(row=row22_idx, column=c)
                    cell.value = None
                # apply styles from row21 if available
                source_row = row22_idx - 1
                if source_row >= 1:
                    for c in range(1, orig_max_col + 1):
                        try:
                            src = ws.cell(row=source_row, column=c)
                            tgt = ws.cell(row=row22_idx, column=c)
                            style = {
                                "font": copy(src.font),
                                "border": copy(src.border),
                                "fill": copy(src.fill),
                                "alignment": copy(src.alignment),
                                "number_format": src.number_format,
                                "protection": copy(src.protection),
                                "_style": copy(getattr(src, "_style", None)),
                            }
                            _apply_style(tgt, style)
                        except Exception:
                            pass
                else:
                    # fallback: set default empty styles
                    for c in range(1, orig_max_col + 1):
                        try:
                            cell = ws.cell(row=row22_idx, column=c)
                            cell.font = Font()
                            cell.fill = PatternFill()
                            cell.border = Border()
                            cell.alignment = Alignment()
                            cell.number_format = 'General'
                        except Exception:
                            pass
            except Exception:
                pass
        # Now write formatting and merges into the newly inserted rows
        for i in range(len(rows)):
            rnum = start_row + i
            # сначала назначаем базовые атрибуты для каждой ячейки строки
            for col in range(1, orig_max_col + 1):
                new_cell = ws.cell(row=rnum, column=col)
                style = col_styles[col]
                # назначаем основные атрибуты стиля
                _apply_style(new_cell, style)
                val = template_values.get(col)
                if val is not None:
                    new_cell.value = copy(val)

            # Восстановить атрибуты RowDimension (без customHeight)
            rd = ws.row_dimensions[rnum]
            if template_row_attrs["height"] is not None:
                rd.height = template_row_attrs["height"]
            rd.hidden = template_row_attrs["hidden"]
            rd.outline_level = template_row_attrs["outline_level"]

            # Воссоздаём объединённые ячейки: сначала снимаем пересекающиеся объединения, очищаем правые и объединяем
            for min_col, max_col in merged_to_copy:
                left_addr = f"{get_column_letter(min_col)}{rnum}"
                right_addr = f"{get_column_letter(max_col)}{rnum}"
                rng = f"{left_addr}:{right_addr}"

                # If the target row already has overlapping merged cells, unmerge them first
                to_unmerge = []
                for ex in list(ws.merged_cells.ranges):
                    emin_col, emin_row, emax_col, emax_row = range_boundaries(str(ex))
                    if emin_row == rnum and not (emax_col < min_col or emin_col > max_col):
                        to_unmerge.append(str(ex))
                for u in to_unmerge:
                    try:
                        ws.unmerge_cells(u)
                    except Exception:
                        pass

                # Очистить правые ячейки в диапазоне (левая оставим)
                for c in range(min_col + 1, max_col + 1):
                    ws.cell(row=rnum, column=c).value = None

                # Создать объединение на целевой строке
                try:
                    ws.merge_cells(rng)
                except Exception:
                    pass

                # Скопировать значение из шаблона (левая ячейка)
                tv = template_values.get(min_col)
                if tv is not None:
                    ws[left_addr] = copy(tv)

                # Повторно применить все стили ко всем ячейкам в объединении (после merge)
                for c in range(min_col, max_col + 1):
                    try:
                        target = ws.cell(row=rnum, column=c)
                        style_c = col_styles.get(c, {})
                        if style_c:
                            _apply_style(target, style_c)
                    except Exception:
                        pass

            # Восстановить ширины столбцов (если нужно)
            for c, w in col_widths.items():
                if w is not None:
                    try:
                        ws.column_dimensions[get_column_letter(c)].width = w
                    except Exception:
                        pass

    # Построим мэппинг колонка->левая колонка объединения
    merged_map = {}
    for min_col, max_col in merged_to_copy:
        for c in range(min_col, max_col + 1):
            merged_map[c] = min_col

    # Записываем строки: в левую ячейку объединения или в указанную колонку
    for idx, row in enumerate(rows):
        rnum = start_row + idx
        for col_letter, val in row.items():
            cidx = column_index_from_string(col_letter)
            if cidx in merged_map:
                left_c = merged_map[cidx]
                ws[f"{get_column_letter(left_c)}{rnum}"] = val
            else:
                ws[f"{col_letter}{rnum}"] = val

    # Final pass: reapply saved column styles to all inserted rows so they visually match template
    for i in range(len(rows)):
        rnum = start_row + i
        for col in range(1, orig_max_col + 1):
            try:
                target = ws.cell(row=rnum, column=col)
                style = col_styles.get(col)
                if style:
                    _apply_style(target, style)
            except Exception:
                pass

    # Auto-fit row heights for inserted rows based on content and column widths
    # We enable wrap_text and estimate required height using font size and chars per line approximation.
    for i in range(len(rows)):
        rnum = start_row + i
        max_height = 0
        for col in range(1, orig_max_col + 1):
            cell = ws.cell(row=rnum, column=col)
            val = cell.value
            if val is None:
                continue
            # ensure text wraps so we can calculate height
            try:
                if getattr(cell, 'alignment', None) is None:
                    cell.alignment = Alignment(wrap_text=True)
                else:
                    # copy alignment to avoid mutating shared style
                    a = copy(cell.alignment)
                    a.wrap_text = True
                    cell.alignment = a
            except Exception:
                pass

            text = str(val)
            # split lines by explicit newlines, then estimate wrapped lines per segment
            parts = text.split('\n')
            total_lines = 0
            for part in parts:
                part_len = len(part)
                # get column width in characters; fall back to 10 if not set
                try:
                    col_w = ws.column_dimensions[get_column_letter(col)].width
                    if col_w is None:
                        chars_per_line = 10
                    else:
                        # Excel column width roughly equals number of '0' chars; use it directly
                        chars_per_line = max(1, int(col_w))
                except Exception:
                    chars_per_line = 10

                # estimate wrapped lines
                wrapped = (part_len + chars_per_line - 1) // chars_per_line if part_len > 0 else 1
                total_lines += max(1, wrapped)

            # determine font size (points); default to 11
            try:
                fsize = getattr(cell.font, 'size', None) or 11
            except Exception:
                fsize = 11

            # approximate line height multiplier (points per line)
            line_height = float(fsize) * 1.25
            height = total_lines * line_height
            if height > max_height:
                max_height = height

        if max_height > 0:
            try:
                ws.row_dimensions[rnum].height = max_height
            except Exception:
                pass

    # Replace the row immediately after the inserted block with a fresh empty row
    next_row = start_row + len(rows)
    try:
        if next_row <= ws.max_row:
            ws.delete_rows(next_row, 1)
        # insert an empty row at the same position
        ws.insert_rows(next_row, 1)
    except Exception:
        pass
    shared_dir = Path("/app/shared")
    wb.save(f"{shared_dir}/{output_path}")

def _apply_style(target, style: dict):
    """Apply saved style dict to target cell robustly."""
    try:
        if style.get("font") is not None:
            target.font = copy(style["font"])
        if style.get("fill") is not None:
            target.fill = copy(style["fill"])
        if style.get("border") is not None:
            target.border = copy(style["border"])
        if style.get("alignment") is not None:
            target.alignment = copy(style["alignment"])
        if style.get("number_format") is not None:
            target.number_format = style["number_format"]
        if style.get("protection") is not None:
            target.protection = copy(style["protection"])
        try:
            if style.get("_style") is not None:
                target._style = copy(style["_style"])
        except Exception:
            # some openpyxl internals may not allow setting _style
            pass
    except Exception:
        pass




@app.post("/generate-multiple")
async def generate_insert_doc(req: RequestMultipleInsert):
    try:
        print(f"Получен insert-запрос: {req}")

        filename = f"zayavka_{req.number}_{datetime.now().year}.xlsx"
        # template is expected in repo root
        repo_root = Path(".")
        template_path = repo_root / "templateV1.xlsx"

        print(f"Используется шаблон: {template_path}")

        if not template_path.exists():
            return JSONResponse({"error": f"Template file {template_path} not found"}, status_code=500)

        # prepare rows_data list of dicts mapping columns to values
        rows_data = []
        for r in req.rows:
            rows_data.append({
                "B": r.source_name,
                "D": r.applicant_name,
                "E": r.address,
                "F": r.dogs_count,
                "G": r.behavior,
                "H": r.urgency,
                "I": r.contact_person,
            })

        # call the internal generator
        try:
            generate_document_multiple(str(template_path), filename, rows_data, start_row=req.start_row, header={"B15": f"Заявка № {req.number}", "B12": datetime.now().strftime("%d.%m.%Y")})
        except Exception as e:
            print(f"Ошибка в generate_document_multiple: {e}")
            return JSONResponse({"error": str(e)}, status_code=500)

        shared_dir = Path("/app/shared")
        if (shared_dir / filename).exists():
            return JSONResponse({"status": "ok", "filename": filename}, status_code=200)
        if Path(filename).exists():
            return JSONResponse({"status": "ok", "filename": filename}, status_code=200)

        return JSONResponse({"error": "File not created"}, status_code=500)
    except Exception as e:
        print(f"Ошибка при создании insert-документа: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)




if __name__ == "__main__":
    uvicorn.run("main:app", host="0.0.0.0", port=8001, reload=True)
